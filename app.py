import os
import pandas as pd
import openpyxl
import base64
import resend
import psycopg2
from psycopg2.extras import RealDictCursor
from flask import Flask, request, jsonify, render_template_string
from datetime import datetime
from zoneinfo import ZoneInfo
from apscheduler.schedulers.background import BackgroundScheduler

app = Flask(__name__)
FUSO_BR = ZoneInfo("America/Sao_Paulo")

# ==========================================
# CONFIGURAÇÕES DE E-MAIL E BANCO DE DADOS
# ==========================================
RESEND_API_KEY = os.environ.get("RESEND_API_KEY", "SUA_API_KEY_RESEND_AQUI")
DATABASE_URL = os.environ.get("DATABASE_URL")
EMAIL_DESTINATARIO = "alexdealm@gmail.com"
MODELO_EXCEL = "Sistema_de_Horas_Trabalhadas_DATATEMPO_CP2_conciliacao_automatica (1).xlsx"

resend.api_key = RESEND_API_KEY

def get_db_connection():
    db_url = os.environ.get("DATABASE_URL")
    if db_url:
        # Limpa aspas, espaços e remove parametros incompativeis com psycopg2 (como pgbouncer=true)
        db_url = db_url.strip().strip('"').strip("'")
        if "?pgbouncer=true" in db_url:
            db_url = db_url.replace("?pgbouncer=true", "")
        if "&pgbouncer=true" in db_url:
            db_url = db_url.replace("&pgbouncer=true", "")
            
        conn = psycopg2.connect(db_url)
        return conn
    else:
        # Fallback para testes locais em SQLite
        import sqlite3
        conn = sqlite3.connect('ponto.db')
        return conn

# --- 1. BASE DE DADOS CADASTRAIS (DATATEMPO) ---
PESSOAS = {
    "1001": "Adriana Pereira dos Santos", "1002": "Alex de Almeida Mariano",
    "1003": "Amanda do Carmo Ribeiro", "1004": "Angela Bertoli",
    "1005": "Carolina Fantini Vidigal Diniz Dias", "1006": "Eliane dos Santos Oliveira",
    "1007": "Fernanda Lopes", "1008": "Flávia Reis", "1009": "Junia Maria Santos",
    "1010": "Nataly Tayê", "1011": "Nathalia Arruda", "1012": "Nelma Lucia dos S. B. Brandão",
    "1013": "Rafaela Maria Resende Lara", "1014": "Robert Filipe Orlando de Souza",
    "1015": "Vitor Guilherme Miguel Rocha", "1016": "Ana Luisa Nardin",
    "1017": "Renato Inácio da Silva"
}

PROJETOS = [
    {"codigo": "PRJ-001", "nome": "BC - CONFIANÇA PIX"}, {"codigo": "PRJ-002", "nome": "CFI / ACCION"},
    {"codigo": "PRJ-003", "nome": "SICOOB CREDICOM"}, {"codigo": "PRJ-004", "nome": "OTEMPO"},
    {"codigo": "PRJ-005", "nome": "ESTADUAL - OTEMPO"}, {"codigo": "PRJ-006", "nome": "CDL/BH"},
    {"codigo": "PRJ-007", "nome": "TOOLKIT"}, {"codigo": "PRJ-008", "nome": "CAMPANHA"},
    {"codigo": "PRJ-009", "nome": "ADMINISTRATIVO"}, {"codigo": "PRJ-010", "nome": "OUTROS"},
    {"codigo": "PRJ-011", "nome": "BETIM - TRACKING"}
]

ATIVIDADES = [
    {"codigo": "ATV-001", "nome": "SUPERVISÃO"}, {"codigo": "ATV-002", "nome": "CHECAGEM"},
    {"codigo": "ATV-003", "nome": "CODIFICAÇÃO"}, {"codigo": "ATV-004", "nome": "ANÁLISE"},
    {"codigo": "ATV-005", "nome": "RECRUTAMENTO"}, {"codigo": "ATV-006", "nome": "APLICAÇÃO / ENTREVISTA"},
    {"codigo": "ATV-007", "nome": "TRATAMENTO DE DADOS"}, {"codigo": "ATV-008", "nome": "RELATÓRIO"},
    {"codigo": "ATV-009", "nome": "ADMINISTRATIVO"}, {"codigo": "ATV-010", "nome": "APOIO OPERACIONAL"},
    {"codigo": "ATV-011", "nome": "TRANSCRIÇÃO"}, {"codigo": "ATV-012", "nome": "MODERAÇÃO"}
]

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS batidas (
            batida_id VARCHAR(50) PRIMARY KEY,
            data_hora VARCHAR(50),
            data VARCHAR(20),
            dia_semana VARCHAR(30),
            pessoa_id VARCHAR(20),
            nome VARCHAR(100),
            tipo VARCHAR(20),
            projeto VARCHAR(100),
            atividade VARCHAR(100),
            observacao TEXT,
            origem VARCHAR(50),
            inicio_vinculado VARCHAR(50)
        )
    ''')
    conn.commit()
    cursor.close()
    conn.close()

init_db()

# --- 2. MOTOR DE FECHAMENTO E ENVIO DE E-MAIL ---
def fechar_folha_e_enviar():
    print("[AUTOMAÇÃO] Iniciando fechamento da folha e envio por e-mail...")
    try:
        if not os.path.exists(MODELO_EXCEL):
            msg = f"A planilha modelo '{MODELO_EXCEL}' não foi encontrada no servidor."
            print(f"[AUTOMAÇÃO ERRO] {msg}")
            return False, msg

        df_pessoas = pd.read_excel(MODELO_EXCEL, sheet_name='Cadastro Pessoas', header=3)
        pessoas_map = {}
        for _, row in df_pessoas.dropna(subset=['ID']).iterrows():
            if str(row['Status']).strip().upper() == 'ATIVO':
                pessoas_map[str(int(row['ID']))] = {
                    'base': str(row['Base']).strip() if pd.notna(row['Base']) else 'Hora',
                    'taxa_util': float(row['R$/h útil']) if pd.notna(row['R$/h útil']) else 0.0,
                    'taxa_fds': float(row['R$/h FDS']) if pd.notna(row['R$/h FDS']) else 0.0,
                    'diaria': float(row['R$/diária']) if pd.notna(row['R$/diária']) else 0.0
                }

        conn = get_db_connection()
        df_batidas = pd.read_sql_query("SELECT * FROM batidas ORDER BY data_hora ASC", conn)
        conn.close()

        if df_batidas.empty:
            msg = "Nenhuma batida registrada no banco de dados para processar."
            print(f"[AUTOMAÇÃO AVISO] {msg}")
            return False, msg

        wb = openpyxl.load_workbook(MODELO_EXCEL)
        sheet_batidas = wb['Batidas']
        sheet_registros = wb['Registros']

        for r in range(5, sheet_batidas.max_row + 1):
            for c in range(1, 14): sheet_batidas.cell(row=r, column=c).value = None
        for r in range(5, sheet_registros.max_row + 1):
            for c in range(1, 20): sheet_registros.cell(row=r, column=c).value = None

        linha_batida = 5
        linha_registro = 5
        inicios_abertos = {}

        for _, batida in df_batidas.iterrows():
            b_id = str(batida['batida_id'])
            dt_obj = pd.to_datetime(batida['data_hora'])
            p_id = str(batida['pessoa_id'])
            nome = str(batida['nome'])
            tipo = str(batida['tipo']).strip().upper()
            projeto = str(batida['projeto'])
            atividade = str(batida['atividade'])
            observacao = str(batida.get('observacao', '')) if pd.notna(batida.get('observacao')) else ''
            
            dia_semana_num = dt_obj.weekday()
            dias_str = ["SEGUNDA-FEIRA", "TERÇA-FEIRA", "QUARTA-FEIRA", "QUINTA-FEIRA", "SEXTA-FEIRA", "SÁBADO", "DOMINGO"]
            dia_semana = dias_str[dia_semana_num]
            tipo_dia = "FDS/FERIADO" if dia_semana_num >= 5 else "DIA ÚTIL"
            inicio_vinculado = str(batida.get('inicio_vinculado', '')) if pd.notna(batida.get('inicio_vinculado')) else ''

            sheet_batidas.cell(row=linha_batida, column=1, value=b_id)
            sheet_batidas.cell(row=linha_batida, column=2, value=dt_obj)
            sheet_batidas.cell(row=linha_batida, column=3, value=f'=IF(B{linha_batida}="","",INT(B{linha_batida}))')
            sheet_batidas.cell(row=linha_batida, column=4, value=dia_semana)
            sheet_batidas.cell(row=linha_batida, column=5, value=int(p_id))
            sheet_batidas.cell(row=linha_batida, column=6, value=nome)
            sheet_batidas.cell(row=linha_batida, column=7, value=tipo)
            sheet_batidas.cell(row=linha_batida, column=8, value=projeto)
            sheet_batidas.cell(row=linha_batida, column=9, value=atividade)
            sheet_batidas.cell(row=linha_batida, column=10, value=observacao)
            sheet_batidas.cell(row=linha_batida, column=11, value="Web HTML")
            sheet_batidas.cell(row=linha_batida, column=12, value=inicio_vinculado)
            sheet_batidas.cell(row=linha_batida, column=13, value="INÍCIO" if tipo == "INÍCIO" else "VINCULADO")
            linha_batida += 1

            if tipo == "INÍCIO":
                reg_id = "R" + b_id[1:]
                info_pessoa = pessoas_map.get(p_id, {'base': 'Hora', 'taxa_util': 0.0, 'taxa_fds': 0.0, 'diaria': 0.0})
                taxa = info_pessoa['taxa_util'] if tipo_dia == "DIA ÚTIL" else info_pessoa['taxa_fds']

                sheet_registros.cell(row=linha_registro, column=1, value=reg_id)
                sheet_registros.cell(row=linha_registro, column=2, value=dt_obj.date())
                sheet_registros.cell(row=linha_registro, column=3, value=dia_semana)
                sheet_registros.cell(row=linha_registro, column=4, value=tipo_dia)
                sheet_registros.cell(row=linha_registro, column=5, value=int(p_id))
                sheet_registros.cell(row=linha_registro, column=6, value=nome)
                sheet_registros.cell(row=linha_registro, column=7, value=projeto)
                sheet_registros.cell(row=linha_registro, column=8, value=atividade)
                sheet_registros.cell(row=linha_registro, column=9, value=dt_obj.strftime("%H:%M:%S"))
                sheet_registros.cell(row=linha_registro, column=13, value=info_pessoa['base'])
                sheet_registros.cell(row=linha_registro, column=14, value=taxa)
                sheet_registros.cell(row=linha_registro, column=15, value=info_pessoa['diaria'])
                sheet_registros.cell(row=linha_registro, column=17, value="EM ABERTO")
                sheet_registros.cell(row=linha_registro, column=18, value=1)
                sheet_registros.cell(row=linha_registro, column=19, value=dt_obj.strftime("%Y-%m"))
                inicios_abertos[b_id] = {'linha': linha_registro, 'data_inicio': dt_obj, 'taxa': taxa, 'base': info_pessoa['base'], 'diaria': info_pessoa['diaria']}
                linha_registro += 1

            elif tipo == "FIM":
                chave_inicio = inicio_vinculado
                if not chave_inicio:
                    for k, v in reversed(list(inicios_abertos.items())):
                        if k.endswith(f"-{p_id}"):
                            chave_inicio = k
                            break

                if chave_inicio in inicios_abertos:
                    dados_inicio = inicios_abertos[chave_inicio]
                    linha_alvo = dados_inicio['linha']
                    sheet_registros.cell(row=linha_alvo, column=10, value=dt_obj.strftime("%H:%M:%S"))
                    horas = max(0.0, (dt_obj - dados_inicio['data_inicio']).total_seconds() / 3600.0)
                    sheet_registros.cell(row=linha_alvo, column=11, value=round(horas, 2))
                    sheet_registros.cell(row=linha_alvo, column=12, value=observacao)
                    valor = dados_inicio['diaria'] if dados_inicio['base'].upper() == "DIÁRIA" else (horas * dados_inicio['taxa'])
                    sheet_registros.cell(row=linha_alvo, column=16, value=round(valor, 2))
                    sheet_registros.cell(row=linha_alvo, column=17, value="CONCLUÍDO")
                    del inicios_abertos[chave_inicio]

        timestamp = datetime.now(FUSO_BR).strftime('%Y%m%d_%H%M%S')
        nome_saida = f'Sistema_Conciliacao_Final_{timestamp}.xlsx'
        wb.save(nome_saida)
        print(f"[AUTOMAÇÃO] Planilha {nome_saida} gerada. Enviando via API do Resend...")

        with open(nome_saida, "rb") as f:
            excel_bytes = f.read()
            excel_b64 = base64.b64encode(excel_bytes).decode('utf-8')

        params = {
            "from": "Sistema Ponto CP2 <onboarding@resend.dev>",
            "to": [EMAIL_DESTINATARIO],
            "subject": f"Relatório de Fechamento de Ponto - {datetime.now(FUSO_BR).strftime('%d/%m/%Y')}",
            "html": "<p>Olá,</p><p>Segue em anexo a planilha consolidada do fechamento automático do ponto com os cálculos atualizados.</p><p><i>E-mail enviado automaticamente pelo Sistema DATATEMPO/CP2.</i></p>",
            "attachments": [
                {
                    "filename": nome_saida,
                    "content": excel_b64
                }
            ]
        }

        email_res = resend.Emails.send(params)
        msg_sucesso = f"Planilha gerada ({nome_saida}) e e-mail enviado via Resend (ID: {email_res.get('id', 'ok')}) para {EMAIL_DESTINATARIO}!"
        print(f"[AUTOMAÇÃO SUCESSO] {msg_sucesso}")
        return True, msg_sucesso
        
    except Exception as e:
        msg_erro = f"Falha no processo via API Resend: {str(e)}"
        print(f"[AUTOMAÇÃO ERRO] {msg_erro}")
        return False, msg_erro

scheduler = BackgroundScheduler()
scheduler.add_job(func=fechar_folha_e_enviar, trigger="cron", hour=18, minute=30, timezone=FUSO_BR)
scheduler.start()

# --- 3. INTERFACE HTML ---
HTML = """
<!DOCTYPE html>
<html lang="pt-BR">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Sistema de Batida de Ponto — DATATEMPO/CP2</title>
    <style>
        body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; margin: 0; padding: 20px; background-color: #f0f2f5; color: #333; }
        .container { max-width: 480px; margin: 20px auto; background: #ffffff; padding: 30px; border-radius: 12px; box-shadow: 0 4px 15px rgba(0,0,0,0.08); }
        h2 { text-align: center; color: #1a365d; margin-top: 0; font-size: 1.4em; border-bottom: 2px solid #e2e8f0; padding-bottom: 10px; }
        .form-group { margin-bottom: 18px; }
        label { display: block; margin-bottom: 6px; font-weight: 600; font-size: 0.9em; color: #4a5568; }
        input, select { width: 100%; padding: 11px; box-sizing: border-box; border: 1px solid #cbd5e0; border-radius: 6px; font-size: 0.95em; }
        input:focus, select:focus { outline: none; border-color: #3182ce; box-shadow: 0 0 0 3px rgba(49, 130, 206, 0.15); }
        .btn-group { display: flex; gap: 10px; margin-top: 10px; }
        .btn { flex: 1; padding: 12px; border: none; border-radius: 6px; color: white; cursor: pointer; font-weight: bold; font-size: 1em; transition: background 0.2s; }
        .btn-inicio { background-color: #28a745; }
        .btn-inicio:hover { background-color: #218838; }
        .btn-fim { background-color: #dc3545; }
        .btn-fim:hover { background-color: #c82333; }
        .btn-espelho { background-color: #3182ce; width: 100%; margin-top: 20px; }
        .btn-espelho:hover { background-color: #2b6cb0; }
        .extrato { margin-top: 25px; display: none; background: #faf5ff; padding: 15px; border-radius: 8px; border: 1px solid #e9d8fd; }
        .extrato h3 { margin-top: 0; font-size: 1.05em; color: #553c9a; }
        table { width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 0.85em; }
        th, td { border: 1px solid #e2e8f0; padding: 8px; text-align: left; }
        th { background-color: #edf2f7; color: #2d3748; }
        .badge-inicio { color: #276749; font-weight: bold; }
        .badge-fim { color: #9b2c2c; font-weight: bold; }
    </style>
</head>
<body>
    <div class="container">
        <h2>DATATEMPO / CP2<br><span style="font-size: 0.7em; font-weight: normal; color: #718096;">Registro Individual de Ponto</span></h2>
        
        <div class="form-group">
            <label for="pessoa_select">Selecione seu Nome/ID:</label>
            <select id="pessoa_select" onchange="atualizarId()">
                <option value="">-- Selecione na lista --</option>
                {% for id, nome in pessoas.items() %}
                    <option value="{{ id }}">{{ id }} - {{ nome }}</option>
                {% endfor %}
            </select>
        </div>

        <div class="form-group">
            <label for="pessoa_id">Ou digite diretamente seu ID:</label>
            <input type="text" id="pessoa_id" placeholder="Ex: 1002" oninput="atualizarSelect()">
        </div>

        <div class="form-group">
            <label for="projeto">Projeto:</label>
            <select id="projeto">
                {% for p in projetos %}
                    <option value="{{ p.nome }}">{{ p.codigo }} - {{ p.nome }}</option>
                {% endfor %}
            </select>
        </div>

        <div class="form-group">
            <label for="atividade">Atividade:</label>
            <select id="atividade">
                {% for a in atividades %}
                    <option value="{{ a.nome }}">{{ a.codigo }} - {{ a.nome }}</option>
                {% endfor %}
            </select>
        </div>

        <div class="form-group">
            <label for="observacao">Observação (Opcional):</label>
            <input type="text" id="observacao" placeholder="Ex: Checagem de amostragem no bairro X">
        </div>

        <div class="btn-group">
            <button class="btn btn-inicio" id="btn-in" onclick="registrar('INÍCIO')">INÍCIO</button>
            <button class="btn btn-fim" id="btn-out" onclick="registrar('FIM')">FIM</button>
        </div>

        <button class="btn btn-espelho" onclick="carregarEspelho()">Ver Meu Espelho de Hoje</button>

        <div class="extrato" id="extrato">
            <h3 id="titulo-espelho">Espelho do Ponto</h3>
            <table>
                <thead>
                    <tr>
                        <th>Hora</th>
                        <th>Tipo</th>
                        <th>Projeto</th>
                        <th>Atividade</th>
                    </tr>
                </thead>
                <tbody id="tabela-espelho"></tbody>
            </table>
        </div>
    </div>

    <script>
        function atualizarId() {
            const select = document.getElementById('pessoa_select');
            document.getElementById('pessoa_id').value = select.value;
        }

        function atualizarSelect() {
            const idInput = document.getElementById('pessoa_id').value;
            const select = document.getElementById('pessoa_select');
            select.value = idInput;
        }

        async function registrar(tipo) {
            const pessoa_id = document.getElementById('pessoa_id').value;
            if (!pessoa_id) {
                alert("Por favor, selecione ou digite o seu ID de funcionário.");
                return;
            }

            document.getElementById('btn-in').disabled = true;
            document.getElementById('btn-out').disabled = true;

            const payload = {
                pessoa_id: pessoa_id,
                projeto: document.getElementById('projeto').value,
                atividade: document.getElementById('atividade').value,
                observacao: document.getElementById('observacao').value,
                tipo: tipo
            };

            try {
                const response = await fetch('/api/bater_ponto', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify(payload)
                });

                const result = await response.json();
                if (response.ok) {
                    alert(result.mensagem);
                    document.getElementById('observacao').value = '';
                    carregarEspelho();
                } else {
                    alert("Erro: " + result.erro);
                }
            } catch (err) {
                alert("Erro de conexão ao salvar ponto.");
            } finally {
                document.getElementById('btn-in').disabled = false;
                document.getElementById('btn-out').disabled = false;
            }
        }

        async function carregarEspelho() {
            const pessoa_id = document.getElementById('pessoa_id').value;
            if (!pessoa_id) {
                alert("Selecione seu ID para carregar seu espelho de ponto.");
                return;
            }

            const response = await fetch('/api/espelho?id=' + pessoa_id);
            const dados = await response.json();

            if (!response.ok) {
                alert("Erro ao buscar espelho: " + dados.erro);
                return;
            }

            const tbody = document.getElementById('tabela-espelho');
            tbody.innerHTML = '';

            dados.registros.forEach(b => {
                const tr = document.createElement('tr');
                const hora = b.data_hora.split(' ')[1];
                const badgeClass = b.tipo === 'INÍCIO' ? 'badge-inicio' : 'badge-fim';
                
                tr.innerHTML = `
                    <td>${hora}</td>
                    <td class="${badgeClass}">${b.tipo}</td>
                    <td>${b.projeto}</td>
                    <td>${b.atividade}</td>
                `;
                tbody.appendChild(tr);
            });

            document.getElementById('titulo-espelho').innerText = `Espelho de Hoje - ${dados.nome}`;
            document.getElementById('extrato').style.display = 'block';
        }
    </script>
</body>
</html>
"""

# --- 4. ROTAS DA APLICAÇÃO ---
@app.route('/')
def index():
    return render_template_string(HTML, pessoas=PESSOAS, projetos=PROJETOS, atividades=ATIVIDADES)

@app.route('/api/bater_ponto', methods=['POST'])
def bater_ponto():
    dados = request.json
    pessoa_id = str(dados.get('pessoa_id', '')).strip()
    
    if pessoa_id not in PESSOAS:
        return jsonify({"erro": f"ID {pessoa_id} não encontrado!"}), 400

    nome_pessoa = PESSOAS[pessoa_id]
    agora = datetime.now(FUSO_BR)
    
    data_hora = agora.strftime('%Y-%m-%d %H:%M:%S')
    data = agora.strftime('%Y-%m-%d')
    dias_pt = ['SEGUNDA-FEIRA', 'TERÇA-FEIRA', 'QUARTA-FEIRA', 'QUINTA-FEIRA', 'SEXTA-FEIRA', 'SÁBADO', 'DOMINGO']
    dia_semana = dias_pt[agora.weekday()]
    
    batida_id = f"B{agora.strftime('%Y%m%d%H%M%S')}-{pessoa_id}"
    tipo = dados.get('tipo')
    projeto = dados.get('projeto')
    atividade = dados.get('atividade')
    observacao = dados.get('observacao', '')
    origem = "Web HTML"

    conn = get_db_connection()
    cursor = conn.cursor()
    inicio_vinculado = None
    if tipo == 'FIM':
        if DATABASE_URL:
            cursor.execute('''SELECT batida_id FROM batidas WHERE pessoa_id = %s AND tipo = 'INÍCIO' AND data = %s ORDER BY data_hora DESC LIMIT 1''', (pessoa_id, data))
        else:
            cursor.execute('''SELECT batida_id FROM batidas WHERE pessoa_id = ? AND tipo = 'INÍCIO' AND data = ? ORDER BY data_hora DESC LIMIT 1''', (pessoa_id, data))
        row = cursor.fetchone()
        if row: inicio_vinculado = row[0]

    placeholder = "(%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)" if DATABASE_URL else "(?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
    cursor.execute(f'''INSERT INTO batidas VALUES {placeholder}''', 
                   (batida_id, data_hora, data, dia_semana, pessoa_id, nome_pessoa, tipo, projeto, atividade, observacao, origem, inicio_vinculado))
    conn.commit()
    cursor.close()
    conn.close()
    return jsonify({"mensagem": f"Ponto de {tipo} registrado com sucesso para {nome_pessoa} às {agora.strftime('%H:%M:%S')}!"})

@app.route('/api/espelho', methods=['GET'])
def get_espelho():
    pessoa_id = str(request.args.get('id', '')).strip()
    if pessoa_id not in PESSOAS: return jsonify({"erro": "ID inválido."}), 400
    agora = datetime.now(FUSO_BR)
    data_hoje = agora.strftime('%Y-%m-%d')
    
    conn = get_db_connection()
    if DATABASE_URL:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
        cursor.execute("SELECT data_hora, tipo, projeto, atividade, observacao FROM batidas WHERE pessoa_id = %s AND data = %s ORDER BY data_hora ASC", (pessoa_id, data_hoje))
        registros = [dict(row) for row in cursor.fetchall()]
    else:
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()
        cursor.execute("SELECT data_hora, tipo, projeto, atividade, observacao FROM batidas WHERE pessoa_id = ? AND data = ? ORDER BY data_hora ASC", (pessoa_id, data_hoje))
        registros = [dict(row) for row in cursor.fetchall()]
        
    cursor.close()
    conn.close()
    return jsonify({"nome": PESSOAS[pessoa_id], "registros": registros})

@app.route('/api/todas_batidas', methods=['GET'])
def get_todas_batidas():
    conn = get_db_connection()
    if DATABASE_URL:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
    else:
        cursor = conn.cursor()
    cursor.execute("SELECT * FROM batidas ORDER BY data_hora ASC")
    registros = [dict(row) for row in cursor.fetchall()]
    cursor.close()
    conn.close()
    return jsonify(registros)

@app.route('/admin', methods=['GET'])
def painel_admin():
    conn = get_db_connection()
    if DATABASE_URL:
        cursor = conn.cursor(cursor_factory=RealDictCursor)
    else:
        cursor = conn.cursor()
    cursor.execute("SELECT * FROM batidas ORDER BY data_hora DESC")
    registros = [dict(row) for row in cursor.fetchall()]
    cursor.close()
    conn.close()
    
    html = """
    <!DOCTYPE html>
    <html lang="pt-BR">
    <head>
        <meta charset="UTF-8"><title>Auditoria do Banco - DATATEMPO/CP2</title>
        <style>
            body { font-family: Arial, sans-serif; padding: 20px; }
            table { width: 100%; border-collapse: collapse; margin-top: 20px; font-size: 0.9em; }
            th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
            th { background-color: #3182ce; color: white; }
            tr:nth-child(even) { background-color: #f2f2f2; }
            .badge-inicio { color: green; font-weight: bold; }
            .badge-fim { color: red; font-weight: bold; }
            .btn-limpar { background-color: #dc3545; color: white; border: none; padding: 10px 15px; cursor: pointer; border-radius: 5px; font-weight: bold; }
            .btn-limpar:hover { background-color: #c82333; }
        </style>
    </head>
    <body>
        <h2>Auditoria de Dados na Nuvem (PostgreSQL / Supabase)</h2>
        <div style="display: flex; justify-content: space-between; align-items: center;">
            <p>Total de batidas registradas: <b>{{ registros|length }}</b></p>
            <button class="btn-limpar" onclick="limparBanco()">⚠️ Zerar Banco de Dados (Testes)</button>
        </div>
        <table>
            <tr>
                <th>Data / Hora (BR)</th>
                <th>Pessoa</th>
                <th>Tipo</th>
                <th>Projeto</th>
                <th>ID da Batida</th>
            </tr>
            {% for r in registros %}
            <tr>
                <td>{{ r.data_hora }}</td>
                <td>{{ r.pessoa_id }} - {{ r.nome }}</td>
                <td class="{% if r.tipo == 'INÍCIO' %}badge-inicio{% else %}badge-fim{% endif %}">{{ r.tipo }}</td>
                <td>{{ r.projeto }}</td>
                <td style="font-size: 0.8em; color: gray;">{{ r.batida_id }}</td>
            </tr>
            {% endfor %}
        </table>

        <script>
            function limparBanco() {
                if(confirm("Tem certeza que deseja apagar TODAS as batidas do PostgreSQL?")) {
                    fetch('/api/limpar_banco')
                    .then(r => r.json())
                    .then(d => {
                        alert(d.status);
                        location.reload();
                    });
                }
            }
        </script>
    </body>
    </html>
    """
    return render_template_string(html, registros=registros)

@app.route('/api/limpar_banco', methods=['GET'])
def limpar_banco():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM batidas")
    conn.commit()
    cursor.close()
    conn.close()
    return jsonify({"status": "Banco de dados PostgreSQL limpo com sucesso!"})

@app.route('/api/fechar_agora', methods=['GET'])
def fechar_agora():
    sucesso, mensagem = fechar_folha_e_enviar()
    status_code = 200 if sucesso else 500
    return jsonify({"sucesso": sucesso, "detalhes": mensagem}), status_code

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)